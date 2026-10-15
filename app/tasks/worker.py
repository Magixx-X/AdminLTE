import json
import os
import signal
import threading
import time
from pathlib import Path
from sqlalchemy import select, update
from openpyxl import Workbook
from app.models import QueryTask, ExportRecord, TemplateVersion, Setting, User, now
from app.repositories import audit, limit
from app.security import AppError, cipher, task_access, permissions
from app.providers.base import QueryCancelled, QueryTimeout
from app.tasks.cache import ResultCache, cache_server
from app.services.excel import safe_cell, sheet_name


class Worker:
    def __init__(self, app):
        self.app, self.repo, self.config = app, app.extensions['repo'], app.config
        self.provider = app.extensions['provider']
        self.cache = ResultCache(self.config)
        self.shutdown = threading.Event()
        self.active = {}
        self.lock = threading.Lock()
        self.last_disk_warning = 0

    def recover(self):
        with self.repo.transaction() as db:
            for task in db.scalars(select(QueryTask).where(QueryTask.status == 'RUNNING')):
                task.status, task.error_code, task.error, task.finished_at = 'INTERRUPTED', 'WORKER_RESTART', 'Worker 重启，任务已中断，请重新执行', now()
                audit(db, 'query.interrupted', task.user_id, task_id=task.id, result='INTERRUPTED')
            for record in db.scalars(select(ExportRecord).where(ExportRecord.status == 'RUNNING')):
                record.status, record.error = 'INTERRUPTED', 'Worker 重启，请重新导出'
                audit(db, 'export.interrupted', record.user_id, task_id=record.id, result='INTERRUPTED')
        for partial in Path(self.config['EXPORT_DIR']).glob('*.part'):
            partial.unlink(missing_ok=True)
        for temporary in (Path(self.config['EXPORT_DIR']) / 'temp').glob('openpyxl.*'):
            temporary.unlink(missing_ok=True)

    def claim(self, model):
        with self.repo.transaction() as db:
            db.connection().exec_driver_sql('BEGIN IMMEDIATE')
            item = db.scalar(select(model).where(model.status == 'PENDING').order_by(model.created_at).limit(1))
            if not item:
                return None
            item.status = 'RUNNING'
            if model == QueryTask:
                item.started_at = now()
            audit(db, 'query.running' if model == QueryTask else 'export.running', item.user_id, task_id=item.id)
            return item.id

    def authorized(self, db, task, action):
        user = db.get(User, task.user_id)
        perms = permissions(db, task.user_id)
        needed = 'template:execute' if action == 'execute' else 'result:export'
        if not user or not user.enabled or needed not in perms:
            raise AppError('PERMISSION_REVOKED', '账号或执行权限已撤销', 403)
        task_access(db, task, task.user_id, action)

    def execute_query(self, tid):
        started = time.monotonic()
        try:
            with self.repo.transaction() as db:
                task = db.get(QueryTask, tid)
                self.authorized(db, task, 'execute')
                version = db.get(TemplateVersion, task.version_id)
                params = json.loads(cipher(self.config).decrypt(task.payload.encode()))
                timeout = limit(db, self.config, 'QUERY_TIMEOUT')
                timeout = min(version.scope_policy.get('query_timeout', timeout), self.config['MAX_QUERY_TIMEOUT'])
                cap = limit(db, self.config, 'MAX_PREVIEW_ROWS')
            last_checked, cancelled = [0], [False]
            def stop():
                if self.shutdown.is_set():
                    return True
                if time.monotonic() - last_checked[0] > .1:
                    with self.repo.transaction() as db:
                        current = db.get(QueryTask, tid)
                        cancelled[0] = current.cancel_requested
                        self.authorized(db, current, 'execute')
                    last_checked[0] = time.monotonic()
                return cancelled[0]
            rows, columns, truncated, byte_count = [], [], False, 0
            stream = self.provider.rows(version, params, task.month, task.scopes, stop, timeout)
            try:
                for names, row in stream:
                    if len(rows) >= cap:
                        truncated = True
                        break
                    if len(names) > 200:
                        raise AppError('COLUMN_QUOTA', '列数超过 200，请收窄模板')
                    byte_count += len(json.dumps(row, default=str, ensure_ascii=False).encode())
                    if byte_count > self.config['CACHE_BYTES'] // self.config['QUERY_CONCURRENCY']:
                        raise AppError('CACHE_QUOTA', '单任务预览内存配额已达到，请缩小范围')
                    columns = names
                    rows.append(row)
            finally:
                stream.close()
            if stop():
                raise QueryCancelled()
            self.cache.put(tid, columns, rows)
            with self.repo.transaction() as db:
                item = db.get(QueryTask, tid)
                self.authorized(db, item, 'execute')
                if item.cancel_requested:
                    raise QueryCancelled()
                item.status, item.row_count, item.columns, item.truncated = 'SUCCEEDED', len(rows), columns, truncated
                item.finished_at, item.duration = now(), time.monotonic() - started
                audit(db, 'query.succeeded', item.user_id, task_id=tid, template_id=item.template_id, version_id=item.version_id, row_count=len(rows), duration=item.duration)
        except Exception as exc:
            code, status, message = 'QUERY_FAILED', 'FAILED', '查询失败，请检查审核模板或联系管理员'
            if isinstance(exc, QueryTimeout):
                code, message = 'QUERY_TIMEOUT', '查询超时，请缩小月份或条件范围'
            elif isinstance(exc, QueryCancelled):
                code, status, message = 'CANCELLED', 'CANCELLED', '任务已取消'
                if self.shutdown.is_set():
                    code, status, message = 'WORKER_STOP', 'INTERRUPTED', '服务停止，任务已中断'
            elif isinstance(exc, AppError):
                code, message = exc.code, exc.message
            self.app.logger.warning('Query %s ended: %s (%s)', tid, code, type(exc).__name__)
            with self.repo.transaction() as db:
                item = db.get(QueryTask, tid)
                item.status, item.error_code, item.error = status, code, message
                item.finished_at, item.duration = now(), time.monotonic() - started
                audit(db, 'query.' + status.lower(), item.user_id, task_id=tid, result=status, duration=item.duration)

    def execute_export(self, eid):
        path = Path(self.config['EXPORT_DIR']) / (eid + '.xlsx')
        partial = path.with_suffix('.part')
        book, total = Workbook(write_only=True), 0
        try:
            with self.repo.transaction() as db:
                record = db.get(ExportRecord, eid)
                tasks = db.scalars(select(QueryTask).where(QueryTask.group_id == record.group_id).order_by(QueryTask.month)).all()
                max_rows = limit(db, self.config, 'MAX_EXPORT_ROWS')
                max_size = limit(db, self.config, 'MAX_EXPORT_BYTES')
                timeout = limit(db, self.config, 'QUERY_TIMEOUT')
                if record.expires <= time.time():
                    raise AppError('EXPIRED', '导出任务已过期')
            for task in tasks:
                with self.repo.transaction() as db:
                    self.authorized(db, task, 'export')
                    version = db.get(TemplateVersion, task.version_id)
                    params = json.loads(cipher(self.config).decrypt(task.payload.encode()))
                    task_timeout = min(version.scope_policy.get('query_timeout', timeout), self.config['MAX_QUERY_TIMEOUT'])
                part, sheet_rows, sheet, last_checked = 1, 0, None, [0]
                def stop():
                    if self.shutdown.is_set():
                        return True
                    if time.monotonic() - last_checked[0] > .1:
                        with self.repo.transaction() as db:
                            self.authorized(db, db.get(QueryTask, task.id), 'export')
                        last_checked[0] = time.monotonic()
                    return False
                stream = self.provider.rows(version, params, task.month, task.scopes, stop, task_timeout)
                try:
                    for columns, row in stream:
                        if total >= max_rows:
                            raise AppError('EXPORT_QUOTA', '导出超过行数配额，请缩小范围')
                        if sheet is None or sheet_rows >= self.config['EXCEL_SHEET_ROWS']:
                            if sheet is not None:
                                sheet.close()
                                part += 1
                            sheet = book.create_sheet(sheet_name(task.month, part))
                            sheet.append([safe_cell(sheet, c) for c in columns])
                            sheet_rows = 0
                        sheet.append([safe_cell(sheet, value) for value in row])
                        sheet_rows += 1
                        total += 1
                        if total % 100 == 0:
                            temporary_bytes = sum(Path(ws._writer.out).stat().st_size for ws in book.worksheets if ws._writer)
                            if temporary_bytes > max_size * 4:
                                raise AppError('EXPORT_QUOTA', '导出临时文件超过空间配额')
                    if sheet is None:
                        sheet = book.create_sheet(sheet_name(task.month))
                        sheet.append(['此月份没有符合条件的数据'])
                    sheet.close()
                finally:
                    stream.close()
            book.save(partial)
            if partial.stat().st_size > max_size:
                raise AppError('EXPORT_QUOTA', '导出文件超过大小配额')
            with self.repo.transaction() as db:
                for task in tasks:
                    self.authorized(db, task, 'export')
                if record.expires <= time.time():
                    raise AppError('EXPIRED', '文件生成时已过期')
                partial.replace(path)
                item = db.get(ExportRecord, eid)
                item.status, item.relative_path, item.size, item.row_count = 'SUCCEEDED', path.name, path.stat().st_size, total
                audit(db, 'export.succeeded', item.user_id, task_id=eid, row_count=total)
        except Exception as exc:
            partial.unlink(missing_ok=True)
            path.unlink(missing_ok=True)
            with self.repo.transaction() as db:
                item = db.get(ExportRecord, eid)
                item.status = 'INTERRUPTED' if self.shutdown.is_set() else 'FAILED'
                item.error = exc.message if isinstance(exc, AppError) else '导出失败或超时，请缩小范围并重新执行'
                audit(db, 'export.failed', item.user_id, task_id=eid, result=item.status)
            self.app.logger.warning('Export %s ended: %s', eid, type(exc).__name__)
        finally:
            # openpyxl write_only uses temporary XML: always remove private temp files on failure too.
            for ws in book.worksheets:
                if ws._writer:
                    try:
                        if not ws.closed:
                            ws.close()
                        ws._writer.cleanup()
                    except (OSError, StopIteration):
                        pass
            book.close()

    def maintenance(self):
        import shutil
        with self.repo.transaction() as db:
            disk_warning = limit(db, self.config, 'DISK_WARNING_BYTES')
        if shutil.disk_usage(self.config['EXPORT_DIR']).free < disk_warning and time.time() - self.last_disk_warning > 3600:
            self.app.logger.warning('DISK_SPACE_LOW: 请清理过期导出或扩容受保护数据盘')
            self.last_disk_warning = time.time()
        with self.repo.transaction() as db:
            heartbeat = db.get(Setting, 'worker_heartbeat') or Setting(key='worker_heartbeat')
            heartbeat.value = {'timestamp': time.time(), 'pid': os.getpid(), 'status': 'RUNNING', 'active': len(self.active)}
            db.add(heartbeat)
            for item in db.scalars(select(ExportRecord).where(ExportRecord.expires < time.time(), ExportRecord.status.in_(['SUCCEEDED', 'PENDING']))):
                if item.relative_path:
                    try:
                        (Path(self.config['EXPORT_DIR']) / Path(item.relative_path).name).unlink(missing_ok=True)
                    except PermissionError:
                        self.app.logger.warning('Expired export still in use: %s', item.id)
                        continue
                item.status, item.relative_path = 'EXPIRED', None
            if time.time() - (db.get(Setting, 'last_backup').value if db.get(Setting, 'last_backup') else 0) > 86400:
                from app.services.backup import backup
                backup(self.config['DATA_DIR'], Path(self.config['DATA_DIR']) / 'backups')
                item = db.get(Setting, 'last_backup') or Setting(key='last_backup')
                item.value = time.time()
                db.add(item)
        with self.cache.lock:
            self.cache.purge()

    def tick(self):
        with self.lock:
            self.active = {k: v for k, v in self.active.items() if v[1].is_alive()}
            for model, count, handler in ((QueryTask, self.config['QUERY_CONCURRENCY'], self.execute_query), (ExportRecord, 1, self.execute_export)):
                running = sum(1 for m, _ in self.active.values() if m == model)
                for _ in range(count - running):
                    tid = self.claim(model)
                    if not tid:
                        break
                    thread = threading.Thread(target=handler, args=(tid,), daemon=False)
                    self.active[tid] = (model, thread)
                    thread.start()

    def run(self):
        # ponytail: one Worker supervisor per deployment; SQLite queue persists all pending work.
        lockfile = open(Path(self.config['DATA_DIR']) / 'worker.lock', 'a+b')
        lockfile.seek(0)
        lockfile.write(b'0')
        lockfile.flush()
        lockfile.seek(0)
        try:
            if os.name == 'nt':
                import msvcrt
                msvcrt.locking(lockfile.fileno(), msvcrt.LK_NBLCK, 1)
            else:
                import fcntl
                fcntl.flock(lockfile, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except OSError:
            lockfile.close()
            raise RuntimeError('已有 Worker 运行，禁止启动第二个')
        server = cache_server(self.config, self.cache)
        for sig in (signal.SIGINT, signal.SIGTERM):
            signal.signal(sig, lambda *_: self.shutdown.set())
        self.recover()
        last_maintenance = 0
        try:
            while not self.shutdown.wait(.2):
                self.tick()
                if time.monotonic() - last_maintenance > 5:
                    try:
                        self.maintenance()
                    except Exception as exc:
                        self.app.logger.error('Worker maintenance failed: %s', type(exc).__name__)
                    last_maintenance = time.monotonic()
        finally:
            self.shutdown.set()
            for _, thread in self.active.values():
                thread.join(self.config['MAX_QUERY_TIMEOUT'] + 5)
            server.shutdown()
            server.server_close()
            lockfile.close()
