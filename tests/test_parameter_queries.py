import json
from datetime import datetime
from types import SimpleNamespace
import pytest
from openpyxl import load_workbook
from sqlalchemy import delete
from app.config import ROOT
from app.models import UserDataScope, QueryTask, ExportRecord, Template, TemplateVersion, User
from app.security import AppError, cipher
from app.services.parameters import validate_schema, query_plan, oracle_binds
from conftest import complete

SCHEMA = [{'name': 'start_date', 'type': 'date', 'required': True},
          {'name': 'end_date', 'type': 'date', 'required': True, 'end_after': 'start_date'}]


def test_rounds_template_contract_and_partial_month_split():
    sql = (ROOT / 'docs/sql/rounds_record.sql').read_text(encoding='utf-8')
    spec = json.loads((ROOT / 'docs/sql/rounds_record.json').read_text(encoding='utf-8'))
    validate_schema(spec['schema'], spec['sort_fields'], sql)
    plan = query_plan(SCHEMA, {'start_date': '2023-12-20', 'end_date': '2024-03-04'}, 36)
    assert [month for month, _ in plan] == ['2023-12', '2024-01', '2024-02', '2024-03']
    assert plan[0][1]['start_date'] == '2023-12-20' and plan[-1][1]['end_date'] == '2024-03-04'
    assert plan[2][1]['start_date'] == '2024-02-01' and plan[2][1]['end_date'] == '2024-03-01'
    version = SimpleNamespace(sql=sql, schema=SCHEMA)
    _, binds = oracle_binds(version, plan[2][1], '2024-02')
    assert binds == {'start_date': datetime(2024, 2, 1), 'end_date': datetime(2024, 3, 1)}
    with pytest.raises(AppError):
        query_plan(SCHEMA, {'start_date': '2023-01-01', 'end_date': '2024-03-01'}, 12)


def test_save_query_and_monthly_export_without_review_or_department(app, client):
    with app.extensions['repo'].transaction() as db:
        db.execute(delete(UserDataScope).where(UserDataScope.user_id == 1))
    body = dict(name='日期参数报表', sql='SELECT :start_date AS start_date, :end_date AS end_date FROM dual',
                schema=SCHEMA, sort_fields=[])
    response = client.post('/api/templates', body)
    assert response.status_code == 201, response.json
    template = response.json['data']
    assert template['status'] == 'PUBLISHED' and 'scope_policy' not in template
    group = client.post('/api/queries', {'template_id': template['id'],
            'params': {'start_date': '2024-01-20', 'end_date': '2024-03-04'}}).json['data']
    assert len(group['task_ids']) == 3
    with app.extensions['repo'].transaction() as db:
        tasks = [db.get(QueryTask, tid) for tid in group['task_ids']]
        assert [t.month for t in tasks] == ['2024-01', '2024-02', '2024-03']
        assert json.loads(cipher(app.config).decrypt(tasks[0].payload.encode()))['start_date'] == '2024-01-20'
        assert all(t.scopes == [] for t in tasks)
    complete(app, group)
    export = client.post('/api/exports', {'group_id': group['group_id']}).json['data']
    worker = app.extensions['test_worker']
    assert worker.claim(ExportRecord) == export['id']
    worker.execute_export(export['id'])
    with app.extensions['repo'].transaction() as db:
        record = db.get(ExportRecord, export['id'])
        assert record.status == 'SUCCEEDED'
        book = load_workbook(ROOT / app.config['EXPORT_DIR'] / record.relative_path, read_only=True)
        # Test fixture intentionally splits sheets every ten rows.
        assert all(any(name.startswith(month) for name in book.sheetnames) for month in ['2024-01', '2024-02', '2024-03'])
        book.close()


def test_sql_without_dates_runs_once_and_undefined_parameters_fail(app, client):
    spec = dict(name='单次查询', sql='SELECT :keyword FROM dual', schema=[{'name': 'keyword', 'type': 'string'}], sort_fields=[])
    item = client.post('/api/templates', spec).json['data']
    response = client.post('/api/queries', {'template_id': item['id'], 'params': {'keyword': ''}})
    assert response.status_code == 202 and len(response.json['data']['task_ids']) == 1
    complete(app, response.json['data'])
    assert client.get('/api/queries/' + response.json['data']['task_id']).json['data']['status'] == 'SUCCEEDED'
    assert query_plan([], {}, 36) == [('result', {})]
    with pytest.raises(AppError):
        validate_schema([], [], 'SELECT :missing FROM dual')
    validate_schema([], [], "SELECT ':not_a_bind' FROM dual -- :also_not_a_bind")


def test_datetime_range_and_month_boundary():
    schema = [dict(SCHEMA[0], type='datetime'), dict(SCHEMA[1], type='datetime')]
    plan = query_plan(schema, {'start_date': '2024-02-29T23:00', 'end_date': '2024-03-01T01:00'}, 36)
    assert plan[0][1]['end_date'] == '2024-03-01T00:00'
    assert plan[1][1]['start_date'] == '2024-03-01T00:00'
    assert len(query_plan(SCHEMA, {'start_date': '2024-01-01', 'end_date': '2024-02-01'}, 36)) == 1


def test_multiple_ranges_only_split_selected_range():
    schema = SCHEMA + [{'name': 'admitted_from', 'type': 'date'},
                       {'name': 'admitted_until', 'type': 'date', 'end_after': 'admitted_from'}]
    values = {'start_date': '2024-01-20', 'end_date': '2024-03-04',
              'admitted_from': '2023-01-01', 'admitted_until': '2024-01-01'}
    assert query_plan(schema, values, 36) == [('result', values)]
    schema = [dict(s, split_by_month=True) if s['name'] == 'end_date' else s for s in schema]
    validate_schema(schema, [], 'SELECT :start_date, :end_date, :admitted_from, :admitted_until FROM dual')
    plan = query_plan(schema, values, 36)
    assert [m for m, _ in plan] == ['2024-01', '2024-02', '2024-03']
    assert all(p['admitted_from'] == values['admitted_from'] and p['admitted_until'] == values['admitted_until'] for _, p in plan)
    schema[-1]['split_by_month'] = True
    with pytest.raises(AppError):
        validate_schema(schema, [], 'SELECT 1 FROM dual')


def test_disable_month_split_and_month_typed_range():
    schema = [dict(SCHEMA[0]), dict(SCHEMA[1], split_by_month=False)]
    values = {'start_date': '2024-01-20', 'end_date': '2024-03-04'}
    assert query_plan(schema, values, 36) == [('result', values)]
    schema = [dict(s, type='month') for s in SCHEMA]
    plan = query_plan(schema, {'start_date': '2023-12', 'end_date': '2024-02'}, 36)
    assert [m for m, _ in plan] == ['2023-12', '2024-01']
    assert plan[-1][1]['end_date'] == '2024-02'


def test_upgrade_keeps_accounts_versions_and_enables_saved_templates(app):
    with app.extensions['repo'].transaction() as db:
        original_sql = db.get(TemplateVersion, 1).sql
        original_password = db.get(User, 1).password_hash
        db.get(Template, 1).status = 'REVIEW'
        db.get(Template, 1).published_version = None
        db.connection().exec_driver_sql("UPDATE alembic_version SET version_num='0004'")
    result = app.test_cli_runner().invoke(args=['migrate'])
    assert result.exit_code == 0, result.output
    with app.extensions['repo'].transaction() as db:
        assert db.get(Template, 1).status == 'PUBLISHED' and db.get(Template, 1).published_version == 1
        assert db.get(TemplateVersion, 1).sql == original_sql
        assert db.get(User, 1).password_hash == original_password
