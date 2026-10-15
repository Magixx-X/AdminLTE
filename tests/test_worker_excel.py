import io
import threading
import time
from pathlib import Path
import pytest
from openpyxl import load_workbook, Workbook
from sqlalchemy import select
from app.models import QueryTask, ExportRecord, Setting, UserDataScope
from app.security import AppError
from app.services.excel import safe_cell, sheet_name
from app.services.backup import backup
from app.tasks.cache import cache_server, get_page
from conftest import submit, complete, Client


def test_queue_failure_timeout_cancel_recovery(app,client):
    worker=app.extensions['test_worker']
    group=submit(client,{'month_end':'2024-02','scenario':'failure'})
    complete(app,group)
    assert client.get('/api/queries/'+group['task_id']).json['data']['status']=='FAILED'
    group=submit(client,{'month_end':'2024-02','scenario':'timeout'})
    complete(app,group)
    assert client.get('/api/queries/'+group['task_id']).json['data']['error_code']=='QUERY_TIMEOUT'
    pending=submit(client,{'month_end':'2024-02'})
    assert client.post('/api/queries/'+pending['task_id']+'/cancel').status_code==200
    assert client.get('/api/queries/'+pending['task_id']).json['data']['status']=='CANCELLED'
    running=submit(client,{'month_end':'2024-02','scenario':'slow'})
    assert worker.claim(QueryTask)==running['task_id']
    thread=threading.Thread(target=worker.execute_query,args=(running['task_id'],))
    thread.start()
    time.sleep(.1)
    assert client.post('/api/queries/'+running['task_id']+'/cancel').status_code==200
    thread.join(3)
    assert client.get('/api/queries/'+running['task_id']).json['data']['status']=='CANCELLED'
    interrupted=submit(client,{'month_end':'2024-02'})
    worker.claim(QueryTask)
    worker.recover()
    assert client.get('/api/queries/'+interrupted['task_id']).json['data']['status']=='INTERRUPTED'
    assert client.post('/api/queries/'+interrupted['task_id']+'/retry').status_code==202


def test_concurrency_and_bounded_queue(app,client):
    group=submit(client,{'month_start':'2024-01','month_end':'2024-06','scenario':'slow','sample_rows':15})
    worker=app.extensions['test_worker']
    worker.tick()
    assert len(worker.active)==2
    with app.extensions['repo'].transaction() as db:
        states=[t.status for t in db.scalars(select(QueryTask))]
        assert states.count('RUNNING')==2 and states.count('PENDING')==4
    for _,thread in worker.active.values():
        thread.join(3)
    app.config['QUEUE_LIMIT']=4
    result=client.post('/api/queries',{'template_id':1,'params':{'month_start':'2024-01','month_end':'2024-02'}})
    assert result.status_code==409 and result.json['error']['code']=='QUEUE_FULL'


def test_export_multi_month_split_download_revocation_expiry(app,client):
    group=submit(client)
    complete(app,group)
    record=client.post('/api/exports',{'group_id':group['group_id']}).json['data']
    worker=app.extensions['test_worker']
    assert worker.claim(ExportRecord)==record['id']
    worker.execute_export(record['id'])
    status=client.get('/api/exports/'+record['id']).json['data']
    assert status['status']=='SUCCEEDED',status
    download=client.get('/api/exports/'+record['id']+'/download')
    assert download.status_code==200
    book=load_workbook(io.BytesIO(download.data),read_only=True)
    assert book.sheetnames==['2024-02','2024-02_2','2024-02_3','2024-03','2024-03_2','2024-03_3']
    row=list(book['2024-02'].values)[1]
    assert row[2].startswith('合成样例') and row[5].year==2024
    assert sum(len(list(sheet.values))-1 for sheet in book)==50
    book.close()
    download.close()
    with app.extensions['repo'].transaction() as db:
        db.get(ExportRecord,record['id']).expires=time.time()-1
    assert client.get('/api/exports/'+record['id']+'/download').status_code==410
    worker.maintenance()
    assert not (Path(app.config['EXPORT_DIR'])/(record['id']+'.xlsx')).exists()


def test_export_quota_cleanup_and_delete(app,client):
    group=submit(client,{'month_end':'2024-02'})
    complete(app,group)
    app.config['MAX_EXPORT_ROWS']=5
    record=client.post('/api/exports',{'group_id':group['group_id']}).json['data']
    worker=app.extensions['test_worker']
    worker.claim(ExportRecord)
    worker.execute_export(record['id'])
    assert client.get('/api/exports/'+record['id']).json['data']['status']=='FAILED'
    assert not list(Path(app.config['EXPORT_DIR']).glob('*.part'))
    assert not list((Path(app.config['EXPORT_DIR'])/'temp').glob('openpyxl.*'))
    assert client.request('DELETE','/api/exports/'+record['id']).status_code==200


def test_formula_injection_unicode_dates_null_and_sheet_names():
    from datetime import date
    book=Workbook(write_only=True)
    sheet=book.create_sheet(sheet_name('a/b?:[]*'*8,22))
    assert len(sheet.title)<=31 and '/' not in sheet.title
    values=['=CMD()',' +1','-2','@evil','\t=evil','中文',date(2024,2,29),None,1.5]
    sheet.append([safe_cell(sheet,v) for v in values])
    output=io.BytesIO()
    book.save(output)
    loaded=load_workbook(io.BytesIO(output.getvalue()),read_only=True)
    row=list(loaded.active.values)[0]
    assert all(row[i].startswith("'") for i in range(5))
    assert row[5]=='中文' and row[6].day==29 and row[7] is None
    loaded.close()


def test_preview_limit_and_cache_expiry(app,client):
    app.config['MAX_PREVIEW_ROWS']=10
    group=submit(client)
    complete(app,group)
    task=client.get('/api/queries/'+group['task_id']).json['data']
    assert task['truncated'] and task['row_count']==10
    cache=app.extensions['result_cache']
    old=cache.items[group['task_id']]
    cache.items[group['task_id']]=(0,old[1],old[2],old[3])
    assert client.get('/api/queries/'+group['task_id']+'/results').status_code==410


def test_backup_integrity(app):
    import sqlite3
    path=backup(app.config['DATA_DIR'],Path(app.config['DATA_DIR'])/'backup-test')
    with sqlite3.connect(path) as connection:
        assert connection.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
        assert connection.execute('SELECT COUNT(*) FROM users').fetchone()[0]==1


def test_real_ipc_cache_protocol(app,client):
    import socket
    with socket.socket() as sock:
        sock.bind(('127.0.0.1',0))
        port=sock.getsockname()[1]
    app.config['IPC_PORT']=port
    group=submit(client)
    complete(app,group)
    cache=app.extensions.pop('result_cache')
    server=cache_server(app.config,cache)
    try:
        page=get_page(app,group['task_id'],0,5,'record_id','asc')
        assert len(page['rows'])==5
        assert client.get('/api/queries/'+group['task_id']+'/results?length=5').status_code==200
    finally:
        server.shutdown()
        server.server_close()
        app.extensions['result_cache']=cache


def test_export_concurrency_one(app,client):
    group=submit(client,{'month_end':'2024-02'})
    complete(app,group)
    ids=[client.post('/api/exports',{'group_id':group['group_id']}).json['data']['id'] for _ in range(3)]
    worker=app.extensions['test_worker']
    gate=threading.Event()
    original=worker.execute_export
    def blocked(eid):
        gate.wait(3)
        original(eid)
    worker.execute_export=blocked
    worker.tick()
    assert len(worker.active)==1
    with app.extensions['repo'].transaction() as db:
        assert len(db.scalars(select(ExportRecord).where(ExportRecord.status=='PENDING')).all())==2
    gate.set()
    for _,thread in worker.active.values():
        thread.join(3)


def test_production_starts_without_oracle_and_never_falls_back_to_mock(tmp_path):
    from app import create_app
    config={'APP_ENV':'production','SECRET_KEY':'s'*64,'DATA_DIR':str(tmp_path/'d'),'EXPORT_DIR':str(tmp_path/'e'),'LOG_DIR':str(tmp_path/'l')}
    with pytest.raises(RuntimeError,match='Oracle'):
        create_app(dict(config,DATA_SOURCE='mock'))
    app = create_app(dict(config, DATA_SOURCE='oracle', ORACLE_USER=None))
    assert app.test_cli_runner().invoke(args=['migrate']).exit_code == 0
    from app.security import AppError
    with pytest.raises(AppError) as exc:
        app.extensions['provider'].health()
    assert exc.value.code == 'DATABASE_NOT_CONFIGURED'
    app.extensions['repo'].engine.dispose()
    for handler in list(app.logger.handlers):
        if hasattr(handler, 'baseFilename'):
            handler.close()
            app.logger.removeHandler(handler)
